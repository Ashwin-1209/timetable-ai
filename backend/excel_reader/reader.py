from openpyxl import load_workbook

def read_excel(file_path):

    wb = load_workbook(file_path)

    db = {}

    for sheet in wb.sheetnames:
        rows = list(wb[sheet].iter_rows(values_only = True))

        headers = rows[0]

        sheet_data = []

        for row in rows[1:]:
            data = dict(zip(headers, row))
            sheet_data.append(data)

        db[sheet] = sheet_data

    return db

if __name__ == "__main__":
    file_path = "testcase_1.xlsx"

    print(read_excel(file_path))